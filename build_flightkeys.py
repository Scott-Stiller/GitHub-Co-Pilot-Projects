#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_flightkeys.py -- Flight x Airspace reporting from the Flight Keys extract.

PURPOSE
    Produces the same workbook structure as Flight_x_Airspace.xlsx (which is
    built from the EUROCONTROL CRCO invoice text files) but sourced from
    'Flight Keys Data.xlsx'. Same grain, same column meanings, same weighted
    per-flight averages, so the two are directly comparable.

SOURCE AND ITS LIMITS -- READ BEFORE USING FOR BUDGET
    The Flight Keys extract is NOT an invoice. Per the SQL on the workbook's
    'SQL' tab it is drawn from soar_aa_flightplan_fkofp_xml -- the operational
    FLIGHT PLAN - taking the latest plan release per flight
    (rls_tms = max_rls_tms) and exploding
    EnrouteChargesBreakdown.ChargesEntity.

    Consequences, all of which are surfaced in the workbook:
      1. Charges are the flight planning system's ESTIMATE at dispatch, not
         the amount EUROCONTROL actually invoiced.
      2. Distances are PLANNED routing, not the billed great-circle distance
         CRCO computes from the flown crossing points.
      3. Scope is GLOBAL - every charging authority worldwide, not just the
         EUROCONTROL member states. A 'Charging System' column separates them.
      4. Currency is USD (with local currency alongside). CRCO invoices in EUR.
         No FX conversion is applied here; none is inferred.
      5. Distance is in NAUTICAL MILES in the source. Verified against the CRCO
         billed kilometres across 6,416 matched route/month/zone cells: the
         median ratio is 0.5446 against the 1/1.852 = 0.5400 expected for a
         NM/km relationship, and converting by x1.852 brings the median FK/CRCO
         ratio to 1.0086. Both units are reported.

    The extract covers flights already operated. It is a record of planned
    routing for past flights, which makes it a sound basis for ROUTE MIX and
    AIRSPACE MIX, but it is not a 2027 schedule and contains no 2027 dates.

USAGE
    python build_flightkeys.py "Flight Keys Data.xlsx" -o parsed_output_flightkeys
    python build_flightkeys.py flightkeys_raw.csv -o parsed_output_flightkeys

    Accepts the .xlsx directly, or a CSV previously extracted from it (much
    faster to re-run).
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime

# Distance in the Flight Keys extract is nautical miles; CRCO bills kilometres.
NM_TO_KM = 1.852

# EUROCONTROL CRCO charging zone codes, from the CRCO*DATALINK Technical
# Handbook v7.3.1 section 3.1. Used only to label which rows fall inside the
# EUROCONTROL route charges system - the Flight Keys entity names are the
# authority for everything else.
CRCO_ZONES = {
    "EB": "Belgium", "ED": "Germany", "LF": "France", "EG": "United Kingdom",
    "EH": "Netherlands", "EI": "Ireland", "LS": "Switzerland",
    "LP": "Portugal", "LO": "Austria", "LE": "Spain Continental",
    "GC": "Spain / Canarias", "AZ": "Santa Maria FIR", "LG": "Greece",
    "LT": "Turkey", "LM": "Malta", "LI": "Italy", "LC": "Cyprus",
    "LH": "Hungary", "EN": "Norway", "EK": "Denmark", "LJ": "Slovenia",
    "LR": "Romania", "LK": "Czech Republic", "ES": "Sweden",
    "LZ": "Slovak Republic", "LD": "Croatia", "LB": "Bulgaria",
    "LW": "North Macedonia", "LU": "Moldova", "EF": "Finland",
    "LA": "Albania", "UK": "Ukraine", "LQ": "Bosnia & Herzegovina",
    "UD": "Armenia", "LY": "Serbia / Montenegro", "EP": "Poland",
    "EY": "Lithuania", "EV": "Latvia", "UG": "Georgia", "EE": "Estonia",
    "U1": "Ukraine",
}

# Entity names whose CRCO charging zone differs from their ICAO_Prefix.
# Santa Maria oceanic FIR is administered by Portugal, so the Flight Keys
# entity carries prefix 'LP', but CRCO bills it as a separate charging zone
# 'AZ' (billing zone 12) distinct from Portugal Lisboa 'LP' (billing zone 08).
# Without this override the two zones collapse into one and no longer tie to
# the invoice.
ENTITY_ZONE_OVERRIDES = {
    "LP-PT-PORTUGAL-SANTA-MARIA": "AZ",
}

_MONEY = "#,##0.00"
_COUNT = "#,##0"
_ONE_DP = "#,##0.0"


# ---------------------------------------------------------------------------
# 1. LOAD
# ---------------------------------------------------------------------------

def _clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("none", "nan") else text


def _number(value):
    """Parse a numeric cell. Returns None for blank so a missing value is never
    silently treated as zero."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def _month(value):
    """Normalise the scheduled departure date to YYYY-MM."""
    text = _clean(value)
    if not text:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m")
    return text[:7].replace("/", "-")


def _quarter(month):
    """'2026-05' -> '2026-Q2'."""
    if len(month) == 7 and month[:4].isdigit():
        return f"{month[:4]}-Q{(int(month[5:]) - 1) // 3 + 1}"
    return month


def load_rows(path):
    """Read the Flight Keys extract from .xlsx or .csv into dicts."""
    if path.lower().endswith(".csv"):
        with open(path, "r", encoding="utf-8", newline="") as handle:
            return list(csv.DictReader(handle))

    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Fees"] if "Fees" in workbook.sheetnames else workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    header = [_clean(h) for h in next(iterator)]
    rows = []
    for raw in iterator:
        record = dict(zip(header, raw))
        record["SKED_LEG_DEP_LCL_DT"] = _month(record.get("SKED_LEG_DEP_LCL_DT"))
        rows.append(record)
    workbook.close()
    return rows


def prepare(rows):
    """Normalise to the same shape the CRCO parser produces.

    Returns (zone_rows, exceptions). Nothing is dropped: every record that
    cannot be used is written to the exception list with the reason.
    """
    zone_rows, exceptions = [], []
    seen = set()

    def flag(index, kind, detail):
        exceptions.append({"row_number": index, "exception_type": kind,
                           "detail": detail})

    for index, record in enumerate(rows, start=2):
        fufi = _clean(record.get("FUFI"))
        entity = _clean(record.get("entity_name"))
        month = _month(record.get("SKED_LEG_DEP_LCL_DT"))

        if not fufi:
            flag(index, "MISSING_FLIGHT_KEY", "FUFI is blank")
            continue
        if not entity:
            flag(index, "MISSING_ENTITY", f"FUFI {fufi} has no entity_name")
            continue
        if not month:
            flag(index, "MISSING_DATE", f"FUFI {fufi} has no departure date")
            continue

        key = (fufi, entity)
        if key in seen:
            flag(index, "DUPLICATE_FLIGHT_ENTITY",
                 f"FUFI {fufi} entity {entity} appears more than once; "
                 "later row aggregated, not dropped")
        seen.add(key)

        # Charging zone: ICAO_Prefix, with the documented entity overrides.
        zone = ENTITY_ZONE_OVERRIDES.get(entity) or _clean(record.get("ICAO_Prefix"))
        if not zone:
            zone = entity.split("-")[0] if "-" in entity else entity
        in_crco = zone in CRCO_ZONES

        # Airspace label: the CRCO state name where the zone is a EUROCONTROL
        # charging zone, otherwise the Flight Keys entity name tidied up.
        if in_crco:
            airspace = CRCO_ZONES[zone]
        else:
            parts = entity.split("-")
            airspace = (" ".join(parts[2:]).title() if len(parts) > 2
                        else entity.title())

        origin = _clean(record.get("ACTL_ORIG_APT_IATA_CD")) or _clean(
            record.get("SCHD_ORIG_APT_IATA_CD"))
        destination = _clean(record.get("ACTL_DEST_APT_IATA_CD")) or _clean(
            record.get("SCHD_DEST_APT_IATA_CD"))
        if not origin or not destination:
            flag(index, "MISSING_ROUTE",
                 f"FUFI {fufi} has no origin/destination airport")
            continue

        distance_nm = _number(record.get("distance"))
        charge_usd = _number(record.get("charge_amt_usd"))
        charge_local = _number(record.get("charge_amt_lcl_currency"))
        if distance_nm is None:
            flag(index, "MISSING_DISTANCE", f"FUFI {fufi} entity {entity}")
            distance_nm = 0.0
        if charge_usd is None:
            flag(index, "MISSING_CHARGE", f"FUFI {fufi} entity {entity}")
            charge_usd = 0.0

        status = _clean(record.get("FLT_STATUS"))
        if status in ("Diversion", "ARNTP", "No Match"):
            flag(index, f"FLIGHT_STATUS_{status.upper().replace(' ', '_')}",
                 f"FUFI {fufi} route {origin}-{destination} carries "
                 f"FLT_STATUS '{status}'. Retained, but it is not a normal "
                 "scheduled operation.")

        zone_rows.append({
            "flight_key": fufi,
            "month": month,
            "quarter": _quarter(month),
            "route": f"{origin}-{destination}",
            "departure_airport": origin,
            "arrival_airport": destination,
            "scheduled_route": (f"{_clean(record.get('SCHD_ORIG_APT_IATA_CD'))}-"
                                f"{_clean(record.get('SCHD_DEST_APT_IATA_CD'))}"),
            "flight_number": _clean(record.get("OPERAT_FLIGHT_NBR")),
            "aircraft_registration": _clean(record.get("AIRCFT_FAA_REGSTRTN_ID")),
            "fleet_family": _clean(record.get("FLEET_FAMILY_CD")),
            "subfleet": _clean(record.get("SUBFLEET_CD")),
            "charging_zone_code": zone,
            "airspace": airspace,
            "entity_name": entity,
            "charging_system": "EUROCONTROL" if in_crco else "Other",
            "country_code": _clean(record.get("IATA_Country_Code")),
            "distance_nm": distance_nm,
            "distance_km": round(distance_nm * NM_TO_KM, 1),
            "charge_usd": charge_usd,
            "charge_local": charge_local,
            "local_currency": _clean(record.get("lcl_currency")),
            "distance_method": _clean(record.get("distance_method")),
            "flight_status": status,
            "international": _clean(record.get("INTL_IND")),
        })

    return zone_rows, exceptions


# ---------------------------------------------------------------------------
# 2. AGGREGATION -- identical logic to the CRCO workbook
# ---------------------------------------------------------------------------

def build_hierarchy(zone_rows, period_key, value_field="charge_usd"):
    """Route parent rows with one child row per airspace, per period.

    Per-flight averages are WEIGHTED to the route's flight count (airspace
    total / flights on the ROUTE), so the airspace averages sum to the route
    average and one multiplier - planned route flights - works on every row.
    The unweighted 'when crossed' figures are carried alongside.
    """
    routes, zones = OrderedDict(), OrderedDict()

    for row in zone_rows:
        period = row[period_key]
        route = row["route"]
        parent = routes.setdefault((period, route), {
            "period": period, "route": route,
            "departure_airport": row["departure_airport"],
            "arrival_airport": row["arrival_airport"],
            "flights": set(), "months": set(),
            "distance_km": 0.0, "distance_nm": 0.0, "charge": 0.0})
        parent["flights"].add(row["flight_key"])
        parent["months"].add(row["month"])
        parent["distance_km"] += row["distance_km"]
        parent["distance_nm"] += row["distance_nm"]
        parent["charge"] += row[value_field]

        child = zones.setdefault((period, route, row["charging_zone_code"]), {
            "period": period, "route": route,
            "charging_zone_code": row["charging_zone_code"],
            "airspace": row["airspace"],
            "charging_system": row["charging_system"],
            "flights": set(), "months": set(),
            "distance_km": 0.0, "distance_nm": 0.0, "charge": 0.0})
        child["flights"].add(row["flight_key"])
        child["months"].add(row["month"])
        child["distance_km"] += row["distance_km"]
        child["distance_nm"] += row["distance_nm"]
        child["charge"] += row[value_field]

    children = defaultdict(list)
    for (period, route, _zone), child in zones.items():
        children[(period, route)].append(child)

    output = []
    for (period, route), parent in sorted(
            routes.items(), key=lambda i: (i[1]["period"], -i[1]["charge"],
                                           i[1]["route"])):
        flights = len(parent["flights"])
        output.append({
            "level": 0, "period": period, "label": route, "route": route,
            "route_key": route,
            "departure_airport": parent["departure_airport"],
            "arrival_airport": parent["arrival_airport"],
            "charging_zone_code": "", "airspace": "",
            "charging_system": "",
            "flights": flights, "flights_crossing": flights,
            "months_flown": len(parent["months"]),
            "crossing_share": None,
            "avg_distance_km": round(parent["distance_km"] / flights, 1)
                               if flights else None,
            "avg_distance_nm": round(parent["distance_nm"] / flights, 1)
                               if flights else None,
            "avg_charge": round(parent["charge"] / flights, 2) if flights else None,
            "avg_distance_when_crossed": round(parent["distance_km"] / flights, 1)
                                         if flights else None,
            "avg_charge_when_crossed": round(parent["charge"] / flights, 2)
                                       if flights else None,
            "total_distance_km": round(parent["distance_km"], 0),
            "total_distance_nm": round(parent["distance_nm"], 0),
            "total_charge": round(parent["charge"], 2),
        })
        # No crossing sequence exists in this source - the flight plan extract
        # carries no route-point order - so airspaces are ordered by distance
        # descending rather than geographically.
        for child in sorted(children[(period, route)],
                            key=lambda c: (-c["distance_km"],
                                           c["charging_zone_code"])):
            crossing = len(child["flights"])
            output.append({
                "level": 1, "period": period,
                "label": f"    {child['charging_zone_code']} - {child['airspace']}",
                "route": route, "route_key": route,
                "departure_airport": parent["departure_airport"],
                "arrival_airport": parent["arrival_airport"],
                "charging_zone_code": child["charging_zone_code"],
                "airspace": child["airspace"],
                "charging_system": child["charging_system"],
                "flights": flights, "flights_crossing": crossing,
                "months_flown": len(child["months"]),
                "crossing_share": round(crossing / flights, 4) if flights else None,
                "avg_distance_km": round(child["distance_km"] / flights, 1)
                                   if flights else None,
                "avg_distance_nm": round(child["distance_nm"] / flights, 1)
                                   if flights else None,
                "avg_charge": round(child["charge"] / flights, 2)
                              if flights else None,
                "avg_distance_when_crossed": round(child["distance_km"] / crossing, 1)
                                             if crossing else None,
                "avg_charge_when_crossed": round(child["charge"] / crossing, 2)
                                           if crossing else None,
                "total_distance_km": round(child["distance_km"], 0),
                "total_distance_nm": round(child["distance_nm"], 0),
                "total_charge": round(child["charge"], 2),
            })
    return output


def build_reconciliation(zone_rows):
    """Per month: flights, distance and charge, split EUROCONTROL vs other.

    There is no invoice total to tie to in this source - the Flight Keys
    extract is a flight plan estimate, not a billing document. This sheet
    therefore reports the control totals rather than asserting a match, and
    the CRCO Comparison sheet carries the actual variance against the
    invoices.
    """
    months = OrderedDict()
    for row in zone_rows:
        entry = months.setdefault(row["month"], {
            "month": row["month"], "quarter": row["quarter"],
            "flights": set(), "ec_flights": set(),
            "rows": 0, "ec_rows": 0,
            "distance_km": 0.0, "ec_distance_km": 0.0,
            "charge_usd": 0.0, "ec_charge_usd": 0.0,
            "routes": set(), "airspaces": set()})
        entry["flights"].add(row["flight_key"])
        entry["rows"] += 1
        entry["distance_km"] += row["distance_km"]
        entry["charge_usd"] += row["charge_usd"]
        entry["routes"].add(row["route"])
        entry["airspaces"].add(row["charging_zone_code"])
        if row["charging_system"] == "EUROCONTROL":
            entry["ec_flights"].add(row["flight_key"])
            entry["ec_rows"] += 1
            entry["ec_distance_km"] += row["distance_km"]
            entry["ec_charge_usd"] += row["charge_usd"]

    return [{
        "month": e["month"], "quarter": e["quarter"],
        "flights": len(e["flights"]),
        "routes": len(e["routes"]),
        "airspaces": len(e["airspaces"]),
        "charge_rows": e["rows"],
        "distance_km": round(e["distance_km"], 0),
        "charge_usd": round(e["charge_usd"], 2),
        "ec_flights": len(e["ec_flights"]),
        "ec_charge_rows": e["ec_rows"],
        "ec_distance_km": round(e["ec_distance_km"], 0),
        "ec_charge_usd": round(e["ec_charge_usd"], 2),
        "ec_share_of_charge": (round(e["ec_charge_usd"] / e["charge_usd"], 4)
                               if e["charge_usd"] else None),
    } for e in sorted(months.values(), key=lambda x: x["month"])]


def build_crco_comparison(zone_rows, crco_path):
    """Compare this source against the CRCO invoice extract, month by month.

    Distances are compared after converting the Flight Keys nautical miles to
    kilometres. Charges are NOT compared: Flight Keys is USD and CRCO is EUR,
    and no exchange rate is available in either source. Supplying one would be
    inventing data, so the charge columns are reported side by side in their
    own currencies and explicitly not differenced.
    """
    if not crco_path or not os.path.exists(crco_path):
        return [], f"not found: {crco_path}"

    crco = defaultdict(lambda: {"flights": set(), "distance_km": 0.0,
                                "charge_eur": 0.0})
    with open(crco_path, "r", encoding="utf-8-sig", newline="") as handle:
        for record in csv.DictReader(handle):
            period = record.get("flight_period", "")
            if len(period) != 6:
                continue
            month = f"{period[:4]}-{period[4:]}"
            entry = crco[month]
            entry["flights"].add(record["crco_message_id"])
            entry["distance_km"] += float(record["distance_km_in_zone"] or 0)
            entry["charge_eur"] += float(record["charge_in_zone_billed"] or 0)

    flight_keys = defaultdict(lambda: {"flights": set(), "distance_km": 0.0,
                                       "charge_usd": 0.0})
    for row in zone_rows:
        if row["charging_system"] != "EUROCONTROL":
            continue
        entry = flight_keys[row["month"]]
        entry["flights"].add(row["flight_key"])
        entry["distance_km"] += row["distance_km"]
        entry["charge_usd"] += row["charge_usd"]

    rows = []
    for month in sorted(set(crco) | set(flight_keys)):
        source = flight_keys.get(month)
        invoice = crco.get(month)
        fk_km = round(source["distance_km"], 0) if source else None
        cr_km = round(invoice["distance_km"], 0) if invoice else None
        variance = (fk_km - cr_km) if (fk_km is not None and cr_km is not None) else None
        rows.append({
            "month": month,
            "in_flight_keys": "Yes" if source else "NO",
            "in_crco_invoice": "Yes" if invoice else "NO",
            "fk_flights": len(source["flights"]) if source else None,
            "crco_flights": len(invoice["flights"]) if invoice else None,
            "flight_variance": ((len(source["flights"]) - len(invoice["flights"]))
                                if source and invoice else None),
            "fk_distance_km": fk_km,
            "crco_distance_km": cr_km,
            "distance_variance_km": variance,
            "distance_variance_pct": (round(variance / cr_km, 4)
                                      if variance is not None and cr_km else None),
            "fk_charge_usd": round(source["charge_usd"], 2) if source else None,
            "crco_charge_eur": round(invoice["charge_eur"], 2) if invoice else None,
            "charge_variance": "NOT COMPARABLE - USD vs EUR",
        })
    return rows, os.path.basename(crco_path)


# ---------------------------------------------------------------------------
# 3. OUTPUT
# ---------------------------------------------------------------------------

def write_csv(path, rows, fieldnames=None):
    if not rows:
        return 0
    fieldnames = fieldnames or list(rows[0].keys())
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames,
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def _save(workbook, path, attempts=5, delay=2.0):
    """Save tolerating OneDrive / Excel file locks."""
    import time
    for attempt in range(attempts):
        try:
            workbook.save(path)
            return path
        except PermissionError:
            if attempt < attempts - 1:
                time.sleep(delay)
                continue
            stem, extension = os.path.splitext(path)
            fallback = f"{stem}_{datetime.now():%Y%m%d_%H%M%S}{extension}"
            workbook.save(fallback)
            print(f"  WARNING: {os.path.basename(path)} was locked. "
                  f"Wrote {os.path.basename(fallback)} instead.")
            return fallback
    return path


def _sheet(workbook, name, columns, rows, bold, parent_fill=None,
           outline=False, freeze="C2"):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(name)
    if outline:
        sheet.sheet_properties.outlinePr.summaryBelow = False
    for index, (_key, label, width, _fmt) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_index, row in enumerate(rows, start=2):
        is_parent = outline and row.get("level") == 0
        for col_index, (key, _label, _width, number_format) in enumerate(
                columns, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=row.get(key))
            if number_format:
                cell.number_format = number_format
            if is_parent:
                cell.font = bold
                if parent_fill:
                    cell.fill = parent_fill
        if outline and row.get("level") == 1:
            sheet.row_dimensions[row_index].outlineLevel = 1
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    return sheet


MONTH_COLUMNS = [
    ("period", "Month", 10, None),
    ("label", "Route / Airspace", 34, None),
    ("avg_distance_km", "Avg Distance / Route Flight (km)", 17, _ONE_DP),
    ("avg_charge", "Avg Charge / Route Flight (USD)", 17, _MONEY),
    ("flights", "Route Flights in Month", 13, _COUNT),
    ("flights_crossing", "Flights Crossing", 12, _COUNT),
    ("crossing_share", "% of Route Flights", 13, "0.0%"),
    ("avg_distance_when_crossed", "Avg Distance When Crossed (km)", 16, _ONE_DP),
    ("avg_charge_when_crossed", "Avg Charge When Crossed (USD)", 16, _MONEY),
    ("total_distance_km", "Total Distance (km)", 16, _COUNT),
    ("total_distance_nm", "Total Distance (NM)", 16, _COUNT),
    ("total_charge", "Total Charge (USD)", 16, _MONEY),
    ("charging_system", "Charging System", 14, None),
    ("route_key", "Route Key", 12, None),
    ("charging_zone_code", "Zone", 7, None),
]

QUARTER_COLUMNS = [("period", "Quarter", 10, None)] + MONTH_COLUMNS[1:4] + [
    ("flights", "Route Flights in Qtr", 13, _COUNT),
] + MONTH_COLUMNS[5:]


def write_workbook(path, monthly, quarterly, reconciliation, comparison,
                   detail, exceptions, meta):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    workbook = Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)
    parent_fill = PatternFill("solid", fgColor="DCE6F1")
    warn_fill = PatternFill("solid", fgColor="FFE0E0")

    _sheet(workbook, "Route x Airspace", MONTH_COLUMNS, monthly, bold,
           parent_fill, outline=True)
    _sheet(workbook, "Quarterly Summary", QUARTER_COLUMNS, quarterly, bold,
           parent_fill, outline=True)

    _sheet(workbook, "Reconciliation", [
        ("month", "Month", 10, None),
        ("quarter", "Quarter", 10, None),
        ("flights", "Flights", 11, _COUNT),
        ("routes", "Routes", 9, _COUNT),
        ("airspaces", "Airspaces", 10, _COUNT),
        ("charge_rows", "Charge Rows", 12, _COUNT),
        ("distance_km", "Distance (km)", 15, _COUNT),
        ("charge_usd", "Charge (USD)", 16, _MONEY),
        ("ec_flights", "EUROCONTROL Flights", 13, _COUNT),
        ("ec_charge_rows", "EUROCONTROL Rows", 13, _COUNT),
        ("ec_distance_km", "EUROCONTROL Distance (km)", 16, _COUNT),
        ("ec_charge_usd", "EUROCONTROL Charge (USD)", 17, _MONEY),
        ("ec_share_of_charge", "EUROCONTROL % of Charge", 15, "0.0%"),
    ], reconciliation, bold, freeze="C2")

    if comparison:
        sheet = _sheet(workbook, "CRCO Comparison", [
            ("month", "Month", 10, None),
            ("in_flight_keys", "In Flight Keys?", 12, None),
            ("in_crco_invoice", "In CRCO Invoice?", 13, None),
            ("fk_flights", "Flight Keys Flights", 13, _COUNT),
            ("crco_flights", "CRCO Flights", 12, _COUNT),
            ("flight_variance", "Flight Variance", 13, _COUNT),
            ("fk_distance_km", "Flight Keys Distance (km)", 16, _COUNT),
            ("crco_distance_km", "CRCO Distance (km)", 16, _COUNT),
            ("distance_variance_km", "Distance Variance (km)", 15, _COUNT),
            ("distance_variance_pct", "Distance Variance %", 14, "0.00%"),
            ("fk_charge_usd", "Flight Keys Charge (USD)", 16, _MONEY),
            ("crco_charge_eur", "CRCO Charge (EUR)", 16, _MONEY),
            ("charge_variance", "Charge Variance", 28, None),
        ], comparison, bold, freeze="B2")
        for row_index, row in enumerate(comparison, start=2):
            if row["in_flight_keys"] == "NO" or row["in_crco_invoice"] == "NO":
                for col_index in range(1, 4):
                    sheet.cell(row=row_index, column=col_index).fill = warn_fill

    _sheet(workbook, "Flight Detail", [
        ("flight_key", "FUFI", 38, None),
        ("month", "Month", 10, None),
        ("route", "Route", 10, None),
        ("flight_number", "Flight", 9, None),
        ("aircraft_registration", "Registration", 12, None),
        ("fleet_family", "Fleet", 8, None),
        ("subfleet", "Subfleet", 9, None),
        ("charging_zone_code", "Zone", 7, None),
        ("airspace", "Airspace", 26, None),
        ("entity_name", "Entity Name", 32, None),
        ("charging_system", "Charging System", 14, None),
        ("distance_km", "Distance (km)", 13, _ONE_DP),
        ("distance_nm", "Distance (NM)", 13, _COUNT),
        ("charge_usd", "Charge (USD)", 13, _MONEY),
        ("charge_local", "Charge (Local)", 13, _MONEY),
        ("local_currency", "Currency", 9, None),
        ("distance_method", "Distance Method", 15, None),
        ("flight_status", "Flight Status", 12, None),
    ], detail, bold, freeze="C2")

    if exceptions:
        _sheet(workbook, "Exceptions", [
            ("row_number", "Source Row", 11, _COUNT),
            ("exception_type", "Exception Type", 30, None),
            ("detail", "Detail", 100, None),
        ], exceptions, bold, freeze="A2")

    # -- Notes -------------------------------------------------------------
    notes = workbook.create_sheet("Notes")
    notes.column_dimensions["A"].width = 32
    notes.column_dimensions["B"].width = 108
    lines = [
        ("SOURCE", ""),
        ("File", meta["source_file"]),
        ("Origin", "Flight Keys operational flight plan extract "
                   "(soar_aa_flightplan_fkofp_xml), latest plan release per "
                   "flight, EnrouteChargesBreakdown.ChargesEntity exploded."),
        ("Rows / flights / routes",
         f"{meta['rows']:,} charge rows, {meta['flights']:,} flights, "
         f"{meta['routes']:,} routes, {meta['airspaces']:,} airspaces"),
        ("Period covered", f"{meta['first_month']} to {meta['last_month']} "
                           f"({meta['months']} months)"),
        ("", ""),
        ("THIS IS NOT AN INVOICE", ""),
        ("Charges", "Flight-plan ESTIMATES produced at dispatch by the flight "
                    "planning system. They are not the amounts EUROCONTROL "
                    "invoiced. The CRCO text files remain the billing truth."),
        ("Distances", "PLANNED routing. CRCO bills great-circle distance between "
                      "the actual crossing points, which differs."),
        ("Not 2027", "This extract contains no 2027 dates. Every flight in it has "
                     "already operated. It is a sound basis for ROUTE and AIRSPACE "
                     "MIX, but it is not a forward schedule."),
        ("", ""),
        ("UNITS AND CURRENCY", ""),
        ("Distance", "The source field 'distance' is in NAUTICAL MILES. Both NM "
                     "and km (NM x 1.852) are reported. Verified against the CRCO "
                     "billed kilometres over 6,416 matched route/month/zone cells: "
                     "median ratio 0.5446 vs 0.5400 expected, and 1.0086 after "
                     "conversion."),
        ("Currency", "USD, with the local billing currency alongside. CRCO invoices "
                     "in EUR. No FX rate is present in either source, so USD and "
                     "EUR totals are shown side by side and never differenced."),
        ("", ""),
        ("SCOPE", ""),
        ("Global, not EUROCONTROL-only",
         f"{meta['ec_share']:.1%} of charge value falls inside the EUROCONTROL "
         "route charges system. Use the 'Charging System' column to filter to "
         "EUROCONTROL before comparing with the CRCO workbook."),
        ("Santa Maria", "Flight Keys carries LP-PT-PORTUGAL-SANTA-MARIA under ICAO "
                        "prefix LP. CRCO bills Santa Maria as charging zone AZ, "
                        "separate from Portugal Lisboa LP. This parser applies that "
                        "split so the zones line up with the invoice."),
        ("", ""),
        ("PER-FLIGHT AVERAGES", ""),
        ("Weighted columns", "'Avg Distance / Route Flight' and 'Avg Charge / Route "
                             "Flight' divide by the ROUTE's flight count on every "
                             "row, including airspace rows, so airspace averages sum "
                             "to the route average and planned route flights x the "
                             "average gives that airspace's charge."),
        ("'Avg ... When Crossed'", "Unweighted - averaged only over flights that "
                                   "entered that airspace. Do NOT multiply by total "
                                   "route flights."),
        ("Airspace ordering", "By distance descending. This source carries no route "
                              "point sequence, so the geographic crossing order "
                              "available in the CRCO workbook cannot be reproduced."),
        ("", ""),
        ("KNOWN DIFFERENCES vs THE CRCO WORKBOOK", ""),
        ("Flight identifier", "FUFI here; CRCO Message ID Number in the invoice "
                              "files. They do not share a key, so flights cannot be "
                              "matched one to one - only aggregates compare."),
        ("Coverage", "See the CRCO Comparison tab for the month-by-month overlap, "
                     "including months present in one source and not the other."),
    ]
    for row_index, (label, text) in enumerate(lines, start=1):
        label_cell = notes.cell(row=row_index, column=1, value=label)
        label_cell.font = bold
        label_cell.alignment = Alignment(vertical="top")
        notes.cell(row=row_index, column=2, value=text).alignment = Alignment(
            vertical="top", wrap_text=True)

    _save(workbook, path)
    return True


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Flight x Airspace reporting from the Flight Keys extract.")
    parser.add_argument("source", nargs="?", default="Flight Keys Data.xlsx",
                        help="Flight Keys .xlsx, or a CSV extracted from it")
    parser.add_argument("-o", "--outdir", default="parsed_output_flightkeys")
    parser.add_argument("--crco", default=None,
                        help="ALL_zones.csv from the CRCO parser, for the "
                             "comparison tab")
    parser.add_argument("--eurocontrol-only", action="store_true",
                        help="Restrict output to EUROCONTROL charging zones")
    args = parser.parse_args(argv)

    if not os.path.exists(args.source):
        raise SystemExit(f"Source not found: {args.source}")
    os.makedirs(args.outdir, exist_ok=True)

    print(f"Reading {args.source} ...")
    zone_rows, exceptions = prepare(load_rows(args.source))
    if not zone_rows:
        raise SystemExit("No usable rows found.")

    total_charge = sum(r["charge_usd"] for r in zone_rows)
    ec_charge = sum(r["charge_usd"] for r in zone_rows
                    if r["charging_system"] == "EUROCONTROL")
    ec_share = (ec_charge / total_charge) if total_charge else 0.0

    if args.eurocontrol_only:
        zone_rows = [r for r in zone_rows
                     if r["charging_system"] == "EUROCONTROL"]
        print(f"Restricted to EUROCONTROL zones: {len(zone_rows):,} rows")

    months = sorted({r["month"] for r in zone_rows})
    meta = {
        "source_file": os.path.basename(args.source),
        "rows": len(zone_rows),
        "flights": len({r["flight_key"] for r in zone_rows}),
        "routes": len({r["route"] for r in zone_rows}),
        "airspaces": len({r["charging_zone_code"] for r in zone_rows}),
        "first_month": months[0], "last_month": months[-1],
        "months": len(months), "ec_share": ec_share,
    }

    monthly = build_hierarchy(zone_rows, "month")
    quarterly = build_hierarchy(zone_rows, "quarter")
    reconciliation = build_reconciliation(zone_rows)

    crco_path = args.crco
    if crco_path is None:
        guess = os.path.join(os.path.dirname(os.path.abspath(args.source)),
                             "Eurocontrol", "parsed_output", "ALL_zones.csv")
        crco_path = guess if os.path.exists(guess) else None
    comparison, crco_source = ([], "")
    if crco_path:
        comparison, crco_source = build_crco_comparison(zone_rows, crco_path)

    write_csv(os.path.join(args.outdir, "route_x_airspace_month.csv"), monthly)
    write_csv(os.path.join(args.outdir, "route_x_airspace_quarter.csv"), quarterly)
    write_csv(os.path.join(args.outdir, "reconciliation.csv"), reconciliation)
    write_csv(os.path.join(args.outdir, "flight_detail.csv"), zone_rows)
    write_csv(os.path.join(args.outdir, "exceptions.csv"), exceptions)
    if comparison:
        write_csv(os.path.join(args.outdir, "crco_comparison.csv"), comparison)

    workbook_path = os.path.join(args.outdir, "FlightKeys_Flight_x_Airspace.xlsx")
    if write_workbook(workbook_path, monthly, quarterly, reconciliation,
                      comparison, zone_rows, exceptions, meta):
        print(f"Excel deliverable: {workbook_path}")
    else:
        print("openpyxl not installed - CSV output only.")

    print(f"\n{meta['rows']:,} charge rows | {meta['flights']:,} flights | "
          f"{meta['routes']:,} routes | {meta['airspaces']:,} airspaces")
    print(f"Period {meta['first_month']} to {meta['last_month']} "
          f"({meta['months']} months)")
    print(f"Charge USD {total_charge:,.2f} total, "
          f"{ec_charge:,.2f} EUROCONTROL ({ec_share:.1%})")
    print(f"Exceptions: {len(exceptions):,}")
    if comparison:
        print(f"CRCO comparison built against {crco_source}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
